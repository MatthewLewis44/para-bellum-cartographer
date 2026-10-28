"""Gate for tools/research_convert.py, the research workbook converter.

    uv run python tests/test_research_convert.py            # run
    uv run python tests/test_research_convert.py --bless    # rewrite expected/ (review the diff!)

Builds a SYNTHETIC workbook in memory from tests/fixtures/research_convert/cases.json
(no real catalogue content is in this repository) and checks:

  1. the base case converts and matches expected/ BYTE FOR BYTE, ledger included;
  2. a re-run against the minted ledger, without --mint, writes identical bytes and
     leaves the ledger untouched (the ids are pinned, D4);
  3. --mint on a new row only APPENDS to the ledger, and no existing id moves;
  4. a rename entry appended by hand keeps a design's id across a spelling change;
  5. properties of the shipped files: no prose anywhere, the fold is a category
     edit (ids unchanged), an interior non-ASCII space survives byte-identical
     (D7 c), the lead reading, atoms by table membership;
  6. every refuse case is refused with EXACTLY its one code and writes nothing;
  7. every X code in convert_codes.json has a refuse case, and every case's code
     is in a register.

Exit status 0 only if every check holds; each failure is listed by name.
"""

from __future__ import annotations

import contextlib
import copy
import io
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import research_convert as rc  # noqa: E402

FIX = ROOT / "tests" / "fixtures" / "research_convert"
EXPECTED = FIX / "expected"
CASES = json.loads((FIX / "cases.json").read_text(encoding="utf-8"))
OUT_FILES = ("research.json", "bureau.json", "vendor.json", "design_model.json")
KEY = {"Categories": "Code", "Technologies": "ID", "Doctrines": "ID", "Bureaus": "ID",
       "Designs": "Design", "Vendors": "Vendor", "Start 1930": "Nation", "Treaties 1930": "Nation"}


def apply(base: dict, patch: list) -> dict:
    case = copy.deepcopy(base)
    case.setdefault("ledger", None)
    for op in patch:
        (verb, arg), = op.items()
        sheets = case["sheets"]
        if verb == "drop_sheet":
            del sheets[arg]
        elif verb == "rename_column":
            sheet, old, new = arg
            sheets[sheet][0] = [new if c == old else c for c in sheets[sheet][0]]
        elif verb in ("set", "drop_row"):
            sheet, key = arg[0], arg[1]
            header = sheets[sheet][0]
            k = header.index(KEY[sheet])
            hits = [r for r in sheets[sheet][1:] if r[k] == key]
            assert len(hits) == 1, f"patch {op}: {len(hits)} rows match"
            if verb == "set":
                hits[0][header.index(arg[2])] = arg[3]
            else:
                sheets[sheet].remove(hits[0])
        elif verb == "add_row":
            sheets[arg[0]].append(arg[1])
        elif verb == "overlay":
            case["overlay"].update(arg)
        elif verb == "ledger":
            case["ledger"] = {"format": rc.LEDGER_FORMAT, "designs": [], "vendors": [], **arg}
        elif verb == "map_inputs":   # null: the three map inputs are not written
            case["provinces"] = case["facilities"] = case["map"] = arg
        elif verb == "effects":   # replaces the effects map in the effects file
            case["effects"]["effects"] = arg
        elif verb == "effects_file":   # null: no effects file is written
            case["effects"] = arg
        elif verb == "models":   # null: no template model file is written
            case["models"] = arg
        elif verb == "vocabulary":   # null: no vocabulary file is written
            case["vocabulary"] = arg
        else:
            raise AssertionError(f"unknown patch verb {verb!r}")
    return case


def build(case: dict, d: Path) -> tuple[Path, Path, Path]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in case["sheets"].items():
        ws = wb.create_sheet(name)
        ws.append([name])
        ws.append(["synthetic test sheet"])
        ws.append([])
        for r in rows:
            ws.append(r)
    reg = wb.create_sheet("Registries")
    for r in case["registries"]:
        reg.append(r)
    paths = (d / "wb.xlsx", d / "overlay.json", d / "ledger.json")
    wb.save(paths[0])
    paths[1].write_text(json.dumps(case["overlay"], ensure_ascii=False), encoding="utf-8")
    if case.get("ledger") is not None:
        paths[2].write_text(json.dumps(case["ledger"], ensure_ascii=False), encoding="utf-8")
    if case.get("vocabulary") is not None:
        (d / "vocabulary.json").write_text(json.dumps(case["vocabulary"]), encoding="utf-8")
    if case.get("provinces") is not None:
        def square(pid, x, y):
            ring = [[x, y], [x + 1, y], [x + 1, y + 1], [x, y + 1], [x, y]]
            return {"type": "Feature", "properties": {"province_id": pid},
                    "geometry": {"type": "Polygon", "coordinates": [ring]}}
        (d / "provinces.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": [
            square(pid, x, y) for pid, x, y in case["provinces"]]}), encoding="utf-8")
        rows = ["name,kind,tier,country,lon,lat,deposit,confidence,note"] + [
            f"{n},{k},1,{c},{lon},{lat},,site,synthetic" for n, k, c, lon, lat in case["facilities"]]
        (d / "facilities.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
        (d / "map.json").write_text(json.dumps({"hexes": [
            {"political": {"province_at_start": pid}} for pid in case["map"] + [""]]}), encoding="utf-8")
    if case.get("effects") is not None:
        (d / "effects.json").write_text(json.dumps(case["effects"], ensure_ascii=False), encoding="utf-8")
    if case.get("models") is not None:
        (d / "model.json").write_text(json.dumps(case["models"]), encoding="utf-8")
    return paths


def run(wb: Path, ov: Path, led: Path, out: Path, mint: bool = True) -> tuple[int, str]:
    """The vocabulary and template models are read from beside the workbook, where build() writes them."""
    argv = ["--workbook", str(wb), "--vocabulary", str(wb.parent / "vocabulary.json"),
            "--models", str(wb.parent / "model.json"), "--effects", str(wb.parent / "effects.json"),
            "--provinces", str(wb.parent / "provinces.geojson"), "--facilities", str(wb.parent / "facilities.csv"),
            "--map", str(wb.parent / "map.json"),
            "--overlay", str(ov),
            "--ledger", str(led), "--out", str(out)]
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = rc.main(argv + (["--mint"] if mint else []))
    return code, buf.getvalue()


def refused_codes(text: str) -> set[str]:
    return set(re.findall(r"PB-([GX][0-9]{3}) ", text))


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    bless = "--bless" in sys.argv
    fails: list[str] = []
    passed = 0

    def check(name: str, ok: bool, detail: str = "") -> None:
        nonlocal passed
        if ok:
            passed += 1
        else:
            fails.append(f"{name}: {detail}")

    base = apply(CASES["base"], [])
    with tempfile.TemporaryDirectory() as t:
        d = Path(t)
        wb, ov, led = build(base, d)
        out = d / "out"
        code, text = run(wb, ov, led, out)
        check("accept/base converts", code == 0, text)
        if code != 0:   # nothing downstream can run; name the failure so a mutation's kill is attributable
            print("research_convert: base refused; stopping")
            for f in fails:
                print(f"  FAIL  {f}")
            return 1
        produced = {n: (out / n).read_bytes() for n in OUT_FILES} | {"ledger.json": led.read_bytes()}
        if bless:
            EXPECTED.mkdir(parents=True, exist_ok=True)
            for n, b in produced.items():
                (EXPECTED / n).write_bytes(b)
            print(f"blessed {len(produced)} files into {EXPECTED.relative_to(ROOT)}")
        for n, b in produced.items():
            exp = EXPECTED / n
            check(f"accept/expected {n}", exp.exists() and exp.read_bytes() == b,
                  "differs from expected/ (run --bless only after reviewing)")

        # 2. pinned: re-run on the minted ledger, without --mint
        out2 = d / "out2"
        ledger_before = led.read_bytes()
        code, text = run(wb, ov, led, out2, mint=False)
        check("ledger/rerun without mint converts", code == 0, text)
        for n in OUT_FILES:
            check(f"ledger/rerun identical {n}", code == 0 and (out2 / n).read_bytes() == produced[n])
        check("ledger/rerun leaves ledger untouched", led.read_bytes() == ledger_before)

        research = json.loads(produced["research.json"])
        ids_before = {p["name"]: p["id"] for p in research["projects"] if p["kind"] == "design"}

        # 3. append-only mint
        grown = apply(CASES["base"], [{"add_row": ["Designs", ["Midland", "Armour", "Light tank", "Wren A",
                                                              None, 1932, "import:Oakfield-Marsh", None, None,
                                                              None, None, "COND", "prose"]]}])
        d3 = d / "grow"
        d3.mkdir()
        wb3, ov3, led3 = build(grown, d3)
        led3.write_bytes(ledger_before)
        code, text = run(wb3, ov3, led3, d3 / "out")
        check("ledger/mint converts", code == 0, text)
        if code == 0:
            new = json.loads(led3.read_text(encoding="utf-8"))
            old = json.loads(ledger_before)
            check("ledger/mint appends only", new["designs"][:len(old["designs"])] == old["designs"]
                  and len(new["designs"]) == len(old["designs"]) + 1 and new["vendors"] == old["vendors"])
            after = {p["name"]: p["id"] for p in json.loads((d3 / "out" / "research.json").read_text(
                encoding="utf-8"))["projects"] if p["kind"] == "design"}
            check("ledger/mint moves no existing id", all(after[k] == v for k, v in ids_before.items()))

        # 4. a rename keeps its id through an appended ledger entry
        renamed = apply(CASES["base"], [{"set": ["Designs", "ZX-3", "Design", "ZX-3 Ausf. A"]}])
        d4 = d / "rename"
        d4.mkdir()
        wb4, ov4, led4 = build(renamed, d4)
        led_doc = json.loads(ledger_before)
        led_doc["designs"].append({"id": ids_before["ZX-3"], "name": "ZX-3 Ausf. A"})
        led4.write_text(json.dumps(led_doc, ensure_ascii=False), encoding="utf-8")
        code, text = run(wb4, ov4, led4, d4 / "out", mint=False)
        check("ledger/rename converts without mint", code == 0, text)
        if code == 0:
            after = {p["name"]: p["id"] for p in json.loads((d4 / "out" / "research.json").read_text(
                encoding="utf-8"))["projects"] if p["kind"] == "design"}
            check("ledger/rename keeps the id", after.get("ZX-3 Ausf. A") == ids_before["ZX-3"])

        # 4b. a base-year design in the slice: modelled with year null, year_declared "base",
        #     never a stamped year, because a model's year is HISTORICAL and ranks equipment (D30)
        based = apply(CASES["base"], [{"set": ["Designs", "Lancer Mk I (prototype B)", "Year", "base"]}])
        d5 = d / "based"
        d5.mkdir()
        wb5, ov5, led5 = build(based, d5)
        led5.write_bytes(ledger_before)
        code, text = run(wb5, ov5, led5, d5 / "out", mint=False)
        check("models/a base design in the slice converts", code == 0, text)
        if code == 0:
            m = {x["name"]: x for x in json.loads((d5 / "out" / "design_model.json").read_text(
                encoding="utf-8"))["models"]}["Lancer Mk I (prototype B)"]
            check("models/a base design ships year null, year_declared base, like its project",
                  (m["year"], m["year_declared"]) == (None, "base") and "year_source" not in m, json.dumps(m))

        # 4c. lineage broken at the slice edge: shipped as null and reported BY NAME (D29)
        edge = apply(CASES["base"], [{"set": ["Designs", "Lancer Mk I (prototype B)", "Design bureau", "Kessel"]}])
        d6 = d / "edge"
        d6.mkdir()
        wb6, ov6, led6 = build(edge, d6)
        led6.write_bytes(ledger_before)
        code, text = run(wb6, ov6, led6, d6 / "out", mint=False)
        check("models/a broken chain converts", code == 0, text)
        if code == 0:
            ms = {x["name"]: x for x in json.loads((d6 / "out" / "design_model.json").read_text(
                encoding="utf-8"))["models"]}
            check("models/a parent outside the slice ships supersedes null",
                  sorted(ms) == ["Lancer Mk II"] and ms["Lancer Mk II"]["supersedes"] is None, json.dumps(ms))
            check("models/the broken chain is reported by name",
                  "'Lancer Mk II' -> 'Lancer Mk I (prototype B)': parent is not in the slice" in text, text)

        # 5. properties of the shipped files
        shipped = b"".join(produced[n] for n in OUT_FILES).decode("utf-8")
        check("ships/no prose", "prose" not in shipped, "a prose cell reached the shipped JSON")
        projects = {p.get("authored_id", p["name"]): p for p in research["projects"]}
        sss = projects["SSS-01"]
        check("ships/fold edits the category only", sss["id"] == "pb:sss_01" and sss["category"] == "pb:bbb"
              and sss["authored_id"] == "SSS-01", json.dumps(sss))
        cats = {c["authored_id"]: c for c in research["categories"]}
        check("ships/folded category kept and marked", cats["SSS"]["folded_into"] == "pb:bbb"
              and cats["CCC"]["fallback_into"] == "pb:aaa" and "Total" not in cats)
        nbsp = "Varenne Q-5"
        check("ships/interior non-ASCII space byte-identical (D7 c)", nbsp in projects
              and nbsp.encode() in produced["ledger.json"] and nbsp.encode() in produced["research.json"])
        check("ships/lead is the first own-nation bureau", projects["Lancer Mk I (prototype B)"]["lead"]
              == "pb:de_ostwerk")
        check("ships/lead is null when the first route is an acquisition",
              projects["Kestrel 12(t)"]["lead"] is None)
        kinds = {a["value"]: a["kind"] for p in (sss, projects["AAA-02"]) for c in p["prerequisites"]
                 for a in c if a["kind"] != "xp"}
        check("ships/atoms by table membership, never by prefix", kinds.get("pb:aaa_02") == "technology"
              and kinds.get("pb:bbb_01") == "technology" and kinds.get("pb:bbb_90") == "doctrine",
              json.dumps(kinds))
        check("ships/year declared, never coerced", projects["AAA-01"]["year"] is None
              and projects["AAA-01"]["year_declared"] == "base")
        cost = {k: (p["cost"], p["cost_source"], p["cost_reason"]) for k, p in projects.items()}
        check("ships/cost by rule per kind", cost["AAA-02"] == ([{"currency": "research", "amount": 100}], "rule", None)
              and cost["QL-03"] == ([{"currency": "doctrine", "amount": 80}], "rule", None)
              and cost["Lancer Mk I (prototype B)"][0] == [{"currency": "research", "amount": 60},
                                                         {"currency": "money", "amount": 40},
                                                         {"currency": "steel", "amount": 20}], json.dumps(cost))
        check("ships/a variant costs the fraction, truncated", cost["Lancer Mk II"] ==
              ([{"currency": "research", "amount": 30}, {"currency": "money", "amount": 20},
                {"currency": "steel", "amount": 10}], "rule", None), json.dumps(cost["Lancer Mk II"]))
        check("ships/an override carries its source and reason", cost["BBB-01"] ==
              ([{"currency": "research", "amount": 150}], "override", "synthetic: a deliberately dearer row"))
        check("ships/era multiplier constants, as data", research["era_multiplier"] ==
              {"growth_per_year_per_mille": 1100, "cap_per_mille": 2000})
        models = {m["name"]: m for m in json.loads(produced["design_model.json"])["models"]}
        mk1, mk2 = models.get("Lancer Mk I (prototype B)"), models.get("Lancer Mk II")
        check("models/the slice: active own-nation lead and a classed role", sorted(models) ==
              ["Lancer Mk I (prototype B)", "Lancer Mk II"], json.dumps(sorted(models)))
        check("models/model id is the design id, and unlock names its own study", mk1 is not None
              and mk1["id"] == mk1["unlock"] == projects["Lancer Mk I (prototype B)"]["id"])
        check("models/lineage ships as supersedes", mk2 is not None and mk2["supersedes"] == mk1["id"])
        check("models/template cost and production, with provenance", mk1 is not None
              and (mk1["cost"], mk1["production"], mk1["production_template"], mk1["year"], mk1["year_declared"]) ==
              ({"steel": 4, "oil": 1}, {"units_per_factory_per_month": 2}, "pb:test_tank", 1931, None))
        held = {s["nation"]: s["designs"] for s in research["start"]}
        heron, kestrel = projects["Heron 1927"]["id"], projects["Kestrel 12(t)"]["id"]
        check("start/designs: own designs dated base OR before 1930, not PAPER, derived, never read from prose",
              held == {"NTH": [], "STH": [kestrel], "WST": [heron], "EST": [], "UPL": [], "LOW": [], "MID": []},
              json.dumps(held))
        check("effects/a no_target effect ships as declared", projects["BBB-01"]["effect"] ==
              {"no_target": {"needs": "test_writer", "reason": "synthetic: no writer target exists for this yet"}})
        homes = {b["authored_id"]: b["home"] for b in json.loads(produced["bureau.json"])["bureaus"]}
        check("homes/place, coordinates and the DERIVED province ship side by side",
              homes["DE_OSTWERK"] == {"place": "Ostburg", "lon_e6": 500000, "lat_e6": 500000, "province": "NTH_ONE"},
              json.dumps(homes["DE_OSTWERK"]))
        check("homes/a home in an unshipped province keeps its derived province",
              homes["FR_VARENNE"]["province"] == "STH_ONE", json.dumps(homes["FR_VARENNE"]))
        check("homes/an inactive bureau with no home ships null", homes["DE_OSTWERK_PLANT"] is None)
        check("homes/an inactive bureau may carry a home (it contributes nothing)",
              homes["DE_KESSEL"] == {"place": "Kesselstadt", "lon_e6": 1500000, "lat_e6": 500006, "province": "NTH_TWO"})
        ai = research["ai"]
        check("ai/weights for every nation x LIVE category (the folded one excluded)",
              sorted(ai["weights"]) == sorted(["NTH", "STH", "WST", "EST", "UPL", "LOW", "MID"])
              and all(sorted(v) == ["pb:aaa", "pb:bbb", "pb:ccc"] for v in ai["weights"].values()),
              json.dumps(ai["weights"]))
        check("ai/weights DERIVED from the bureau table: 1000 x count / largest, floored",
              ai["weights"]["NTH"] == {"pb:aaa": 1000, "pb:bbb": 500, "pb:ccc": 100}
              and ai["weights"]["STH"] == {"pb:aaa": 100, "pb:bbb": 100, "pb:ccc": 1000}
              and ai["weights"]["MID"] == {"pb:aaa": 100, "pb:bbb": 100, "pb:ccc": 100}, json.dumps(ai["weights"]))
        check("ai/the counts each weight came from ship beside it",
              ai["weight_inputs"]["NTH"] == {"largest": 2, "bureaus": {"pb:aaa": 2, "pb:bbb": 1, "pb:ccc": 0}}
              and ai["weight_inputs"]["MID"] == {"largest": 0, "bureaus": {"pb:aaa": 0, "pb:bbb": 0, "pb:ccc": 0}}
              and ai["weight_rule"] == {"floor_per_mille": 100}, json.dumps(ai["weight_inputs"]))
        prv = {r["province"]: r for r in ai["province_research_value"]}
        check("ai/province value over SHIPPED provinces only, with its components",
              sorted(prv) == ["NTH_ONE", "NTH_TWO"]
              and prv["NTH_ONE"] == {"province": "NTH_ONE", "value": 130,
                                     "components": {"active_bureaus": 1, "mine": 0, "power_plant": 0, "steel_mill": 1}}
              and prv["NTH_TWO"]["value"] == 0, json.dumps(prv))
        reg = research["registries"]
        check("registries/fielded and focus carry display names",
              reg["fielded"] == [{"id": "pb:test_heavy", "authored_id": "test_heavy", "name": "Test heavy unit"}]
              and reg["focus"] == [{"id": "pb:test_plan", "authored_id": "test_plan", "name": "Test plan"}])
        no_route = [x for x in research["restrictions"] if x["nation"] == "MID"]
        check("ships/a restriction with no covert route ships lifted_by null", no_route ==
              [{"nation": "MID", "treaty": "pb:accord_of_trials", "scope": {"category": "pb:aaa"}, "lifted_by": None}],
              json.dumps(no_route))
        reg = research["registries"]
        check("ships/treaty registry, each enforced", [t["id"] for t in reg["treaties"]] ==
              ["pb:pact_of_testing", "pb:accord_of_trials"]
              and {x["treaty"] for x in research["restrictions"]} == {"pb:pact_of_testing", "pb:accord_of_trials"})
        flags = {f["authored_id"]: f for f in reg["flags"]}
        check("ships/every flag named; only the unreferenced one inert",
              all(f["name"] for f in flags.values())
              and [k for k, f in flags.items() if f["inert_reason"]] == ["test_dormant"], json.dumps(flags))
        bureaus = {b["authored_id"]: b for b in json.loads(produced["bureau.json"])["bureaus"]}
        check("ships/roster cut is a flag", len(bureaus) == 4 and bureaus["DE_OSTWERK"]["capacity"] == 2
              and bureaus["DE_KESSEL"]["active"] is False and bureaus["DE_KESSEL"]["capacity"] is None)

    # 6. refusals
    for case in CASES["refuse"]:
        with tempfile.TemporaryDirectory() as t:
            d = Path(t)
            wb, ov, led = build(apply(CASES["base"], case["patch"]), d)
            out = d / "out"
            code, text = run(wb, ov, led, out, mint=not case.get("no_mint"))
            got = refused_codes(text)
            check(f"refuse/{case['id']}", code == 1 and got == {case["code"]},
                  f"expected exactly {case['code']}, got exit {code} codes {sorted(got)}\n{text}")
            check(f"refuse/{case['id']} writes nothing", not out.exists() or not any(out.iterdir()))

    # 7. register coverage
    ids = [c["id"] for c in CASES["refuse"]]
    check("cases/ids unique", len(ids) == len(set(ids)))
    covered = {c["code"] for c in CASES["refuse"]}
    for code in sorted(set(rc.RULES) - rc.RETIRED):
        check(f"cases/{code} has a refuse case", code in covered)
    for code in sorted(rc.RETIRED):   # retired: nothing may raise it, so nothing may expect it
        check(f"cases/{code} is retired and has no case", code not in covered)
        check(f"cases/{code} is raised nowhere in the converter",
              f'"{code}"' not in (ROOT / "tools" / "research_convert.py").read_text(encoding="utf-8"))
    for code in sorted(covered):
        check(f"cases/{code} is registered", code in rc.RULES or code in rc.rg.RULES)

    print(f"research_convert: {passed} checks passed, {len(fails)} failed; "
          f"{len(CASES['refuse'])} refuse cases over {len(covered)} codes "
          f"({len(rc.RULES)} X codes registered)")
    for f in fails:
        print(f"  FAIL  {f}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
