"""Sweep the research grammar parser over every grammar cell of the research workbook.

    uv run python tools/probe_research_grammar_sweep.py <Para_Bellum_Research_Database.xlsx>

A probe, not the converter: it reads the workbook directly, builds the closed
vocabulary from its own sheets (ids, design names, the Registries sheet,
bureaus per nation, vendors), and runs tools/research_grammar.py over the seven
columns the Grammar sheet governs. It reports, with counts:

  - cells per column and how many are empty (no prerequisite);
  - syntax refusals (G1xx) and unresolved tokens (G2xx), each listed;
  - round-trip: render(parse(cell)) must equal the cell byte for byte;
  - joint designs in the Design bureau column under a naive comma split versus
    the top-level split (commas inside parentheses are part of a name).

Exits non-zero on any refusal, unresolved token or round-trip mismatch. The
workbook is not in this repository; pass its path.
"""

from __future__ import annotations

import collections
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import research_grammar as rg  # noqa: E402

HEADER_ROW = 4   # rows 1-3 are a title, a note and a blank on every data sheet

# (sheet, column, context): the columns the Grammar sheet says it governs.
SWEEP = (("Technologies", "Prerequisites", rg.GATE), ("Technologies", "Cost reducers", rg.GATE),
         ("Doctrines", "Prerequisites", rg.GATE),
         ("Designs", "Prerequisites", rg.GATE), ("Designs", "Cost reducers", rg.GATE),
         ("Designs", "Design bureau", rg.ACQUISITION), ("Designs", "Producers", rg.LIST))


def sheet_rows(wb, name: str) -> list[dict]:
    rows = list(wb[name].iter_rows(min_row=HEADER_ROW, values_only=True))
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]


def registries(wb) -> dict[str, set]:
    """The Registries sheet: titled sections, each a header row then one id per row."""
    found = {"flag": set(), "fielded": set(), "focus": set(), "xp": set()}
    section = None
    for r in wb["Registries"].iter_rows(values_only=True):
        a = r[0]
        if a is None:
            continue
        for title, key in (("Flags", "flag"), ("Fielded", "fielded"), ("Focus", "focus"), ("Experience", "xp")):
            if a.startswith(title):
                section = key
                break
        else:
            if a in ("id", "token") or section is None:
                continue
            if section == "xp":   # 'xp:land:1 … xp:land:3' declares a tier range
                m = re.fullmatch(r"xp:(\w+):(\d+) … xp:\1:(\d+)", a)
                if not m:
                    raise SystemExit(f"Registries: cannot read experience row {a!r}")
                found["xp"] |= {(m.group(1), t) for t in range(int(m.group(2)), int(m.group(3)) + 1)}
            else:
                found[section].add(a)
    return found


def main(path: str) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = openpyxl.load_workbook(path)
    data = {s: sheet_rows(wb, s) for s in ("Technologies", "Doctrines", "Designs", "Bureaus", "Vendors")}
    reg_sets = registries(wb)
    bureaus = collections.defaultdict(set)
    for b in data["Bureaus"]:
        bureaus[b["Nation"]].add(b["Bureau"])
    reg = rg.Registry(
        ids=frozenset([r["ID"] for r in data["Technologies"]] + [r["ID"] for r in data["Doctrines"]]),
        designs=frozenset(g["Design"] for g in data["Designs"]),
        flags=frozenset(reg_sets["flag"]), fielded=frozenset(reg_sets["fielded"]),
        focus=frozenset(reg_sets["focus"]), xp=frozenset(reg_sets["xp"]),
        bureaus={n: frozenset(v) for n, v in bureaus.items()},
        vendors=frozenset(v["Vendor"] for v in data["Vendors"]), nations=frozenset(bureaus))
    print(f"workbook: {path}")
    print(f"vocabulary: {len(reg.ids)} ids, {len(reg.designs)} designs, {len(reg.flags)} flags, "
          f"{len(reg.fielded)} fielded, {len(reg.focus)} focus, {len(reg.xp)} xp, "
          f"{sum(map(len, bureaus.values()))} bureaus in {len(bureaus)} nations, {len(reg.vendors)} vendors")

    problems: list[str] = []
    cells = empty = atoms = 0
    kinds: collections.Counter = collections.Counter()
    for sheet, col, ctx in SWEEP:
        n = e = 0
        for i, row in enumerate(data[sheet]):
            cell = row[col]
            n += 1
            e += cell in (None, "")
            where = f"{sheet}!{col} data row {i + 1}"
            try:
                expr = rg.parse(cell, ctx)
            except rg.GrammarError as err:
                problems.append(f"{where}: {err}")
                continue
            if cell not in (None, "") and rg.render(expr) != cell:
                problems.append(f"{where}: round-trip {cell!r} -> {rg.render(expr)!r}")
            for clause in expr:
                for a in clause:
                    kinds[(col, a.kind)] += 1
                    atoms += 1
            problems += [f"{where}: {p}" for p in rg.resolve(expr, ctx, reg, row.get("Nation"))]
        cells += n
        empty += e
        print(f"  {sheet}.{col} [{ctx}]: {n} cells, {e} empty, {n - e} parsed")
    print(f"total: {cells} cells, {empty} empty, {cells - empty} non-empty, {atoms} tokens")
    print("tokens by column and kind: " + ", ".join(f"{c}/{k} {v}" for (c, k), v in sorted(kinds.items())))
    designs = data["Designs"]
    naive = sum(1 for g in designs if "," in g["Design bureau"])
    top = sum(1 for g in designs if len(rg.parse(g["Design bureau"], rg.ACQUISITION)) > 1)
    print(f"Design bureau: {naive} joint designs by naive comma split, {top} by the top-level split")
    print(f"design names equal to an id: {sorted(reg.designs & reg.ids) or 'none'}")
    if problems:
        print(f"FAIL - {len(problems)} problem(s):")
        for p in problems:
            print(f"  {p}")
        return 1
    print(f"PASS - every one of {cells - empty} non-empty cells parses, resolves and renders back exactly")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    sys.exit(main(sys.argv[1]))
